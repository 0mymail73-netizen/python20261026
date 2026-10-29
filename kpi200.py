import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QStatusBar,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)


SCANNER_URL = "https://scanner.tradingview.com/korea/scan"
APP_DIRECTORY = (
    Path(sys.executable).resolve().parent
    if getattr(sys, "frozen", False)
    else Path(__file__).resolve().parent
)
OUTPUT_FILE = APP_DIRECTORY / "kospi200.xlsx"
PAGE_SIZE = 200
TARGET_COUNT = 200
INDEX_SYMBOL = "KRX:KOSPI200"
SCANNER_COLUMNS = [
    "name",
    "description",
    "indexes",
    "market_cap_basic",
    "close",
    "change",
]
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/130.0.0.0 Safari/537.36"
    ),
    "Origin": "https://www.tradingview.com",
    "Referer": "https://www.tradingview.com/",
}


@dataclass(frozen=True)
class Constituent:
    code: str
    name: str
    market_cap: float
    price: float
    change: float


def parse_scanner_page(
    payload: dict[str, Any],
) -> tuple[int, list[Constituent]]:
    if not isinstance(payload, dict):
        raise ValueError("종목 스캐너 응답 형식이 예상과 다릅니다.")
    total_count = payload.get("totalCount")
    rows = payload.get("data")
    if not isinstance(total_count, int) or not isinstance(rows, list):
        raise ValueError("종목 스캐너 응답 형식이 예상과 다릅니다.")

    constituents: list[Constituent] = []
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError("종목 스캐너의 행 형식이 예상과 다릅니다.")
        values = row.get("d")
        if not isinstance(values, list) or len(values) < len(SCANNER_COLUMNS):
            raise ValueError("종목 스캐너의 데이터 열이 예상과 다릅니다.")

        code, name, indexes, market_cap, price, change = values
        if not isinstance(indexes, list) or not any(
            isinstance(index_info, dict)
            and index_info.get("proname") == INDEX_SYMBOL
            for index_info in indexes
        ):
            continue
        if not isinstance(code, str) or not isinstance(name, str):
            raise ValueError("종목코드 또는 종목명이 올바르지 않습니다.")
        if not all(
            isinstance(value, (int, float))
            for value in (market_cap, price, change)
        ):
            raise ValueError(f"{code} 종목의 시가총액/가격/등락률이 없습니다.")

        constituents.append(
            Constituent(
                code=code,
                name=name.strip(),
                market_cap=float(market_cap),
                price=float(price),
                change=float(change),
            )
        )
    return total_count, constituents


def crawl_kospi200(
    progress_callback: Callable[[int, int], None] | None = None,
) -> tuple[list[Constituent], int]:
    constituents_by_code: dict[str, Constituent] = {}
    offset = 0
    total_count: int | None = None
    page_number = 0

    with requests.Session() as session:
        session.headers.update(HEADERS)
        while total_count is None or (
            offset < total_count and len(constituents_by_code) < TARGET_COUNT
        ):
            end = offset + PAGE_SIZE
            payload = {
                "columns": SCANNER_COLUMNS,
                "filter": [],
                "options": {"lang": "ko"},
                "range": [offset, end],
                "sort": {
                    "sortBy": "market_cap_basic",
                    "sortOrder": "desc",
                },
            }
            response = session.post(SCANNER_URL, json=payload, timeout=30)
            response.raise_for_status()
            response_data = response.json()
            page_total, page_constituents = parse_scanner_page(response_data)
            if total_count is None:
                total_count = page_total
            elif page_total != total_count:
                raise RuntimeError("크롤링 중 전체 종목 수가 변경되었습니다.")

            page_rows = response_data.get("data")
            if not isinstance(page_rows, list) or not page_rows:
                raise RuntimeError("종목 스캐너에서 다음 페이지를 가져오지 못했습니다.")

            for constituent in page_constituents:
                constituents_by_code[constituent.code] = constituent

            page_number += 1
            if progress_callback is not None:
                progress_callback(
                    page_number, len(constituents_by_code)
                )
            offset += len(page_rows)

    all_matches = sorted(
        constituents_by_code.values(),
        key=lambda constituent: constituent.market_cap,
        reverse=True,
    )
    if len(all_matches) < TARGET_COUNT:
        raise RuntimeError(
            f"KOSPI 200 구성종목을 {TARGET_COUNT}개 모두 찾지 못했습니다. "
            f"현재 {len(all_matches)}개를 찾았습니다."
        )
    return all_matches[:TARGET_COUNT], len(all_matches)


def save_to_excel(
    constituents: list[Constituent],
    output_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        raise RuntimeError("엑셀 워크시트를 생성하지 못했습니다.")

    worksheet.title = "KOSPI200"
    headers = ["종목코드", "종목명", "시가총액(원)", "현재가(원)", "등락률"]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="153D35")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for constituent in constituents:
        worksheet.append(
            [
                constituent.code,
                constituent.name,
                constituent.market_cap,
                constituent.price,
                constituent.change / 100,
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = '#,##0;[Red]-#,##0'
    for cell in worksheet["E"][1:]:
        cell.number_format = "0.00%;[Red]-0.00%"
    for column, width in {"A": 14, "B": 32, "C": 20, "D": 16, "E": 12}.items():
        worksheet.column_dimensions[column].width = width
    workbook.save(output_path)


class CrawlerWorker(QThread):
    progress = pyqtSignal(int, int)
    completed = pyqtSignal(object, int)
    failed = pyqtSignal(str)

    def run(self) -> None:
        try:
            constituents, match_count = crawl_kospi200(
                lambda page, count: self.progress.emit(page, count)
            )
        except (
            requests.RequestException,
            ValueError,
            RuntimeError,
            TypeError,
        ) as error:
            self.failed.emit(str(error))
            return
        self.completed.emit(constituents, match_count)


class KOSPI200Window(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.constituents: list[Constituent] = []
        self.worker: CrawlerWorker | None = None
        self.setWindowTitle("KOSPI 200 종목 크롤러")
        self.resize(1000, 700)
        self.status_bar = QStatusBar(self)
        self.setStatusBar(self.status_bar)
        self.setStyleSheet(
            """
            QMainWindow, QWidget#mainPanel {
                background-color: #fff3a6;
            }
            QWidget#heroPanel {
                background-color: #ffd747;
                border: 3px solid #25231e;
                border-radius: 18px;
            }
            QLabel#pageTitle {
                color: #25231e;
                font-size: 25px;
                font-weight: 900;
            }
            QLabel#pageSubtitle {
                color: #4b4638;
                font-size: 13px;
            }
            QWidget#controlCard, QWidget#searchCard, QWidget#tableCard {
                background-color: #fffdf4;
                border: 2px solid #25231e;
                border-radius: 14px;
            }
            QLabel#sectionTitle {
                color: #25231e;
                font-size: 15px;
                font-weight: 800;
                border: none;
            }
            QLabel#resultCount {
                color: #25231e;
                background-color: #b8e9f4;
                border: 2px solid #25231e;
                border-radius: 10px;
                padding: 6px 12px;
                font-weight: 800;
            }
            QLabel#progressLabel {
                color: #343128;
                font-weight: 700;
            }
            QLineEdit {
                color: #25231e;
                background-color: #ffffff;
                border: 2px solid #25231e;
                border-radius: 9px;
                padding: 9px 12px;
                selection-background-color: #44b6d0;
                min-height: 20px;
            }
            QLineEdit:focus {
                background-color: #f2fcff;
                border: 2px solid #2086a0;
            }
            QPushButton {
                color: #25231e;
                background-color: #ffcf38;
                border: 2px solid #25231e;
                border-radius: 9px;
                padding: 9px 16px;
                font-weight: 800;
                min-width: 90px;
            }
            QPushButton:hover {
                background-color: #ffe477;
            }
            QPushButton:pressed {
                background-color: #f0b900;
                padding-top: 11px;
                padding-bottom: 7px;
            }
            QPushButton:disabled {
                color: #777366;
                background-color: #e7e2ce;
                border-color: #918d80;
            }
            QPushButton#crawlButton {
                background-color: #79d5e9;
            }
            QPushButton#crawlButton:hover {
                background-color: #a3e8f4;
            }
            QPushButton#saveButton {
                background-color: #a9dc78;
            }
            QPushButton#saveButton:hover {
                background-color: #c5ee9e;
            }
            QPushButton#clearSearchButton {
                background-color: #f4a1a1;
            }
            QPushButton#clearSearchButton:hover {
                background-color: #ffc0b8;
            }
            QProgressBar {
                color: #25231e;
                background-color: #ffffff;
                border: 2px solid #25231e;
                border-radius: 8px;
                text-align: center;
                font-weight: 800;
                min-height: 20px;
            }
            QProgressBar::chunk {
                background-color: #7dcc76;
                border-radius: 5px;
            }
            QTableWidget {
                color: #25231e;
                background-color: #ffffff;
                alternate-background-color: #fff9dc;
                border: 2px solid #25231e;
                border-radius: 8px;
                gridline-color: #ded7bc;
                selection-background-color: #b8e9f4;
                selection-color: #25231e;
                outline: none;
            }
            QHeaderView::section {
                color: #25231e;
                background-color: #79d5e9;
                border: none;
                border-bottom: 2px solid #25231e;
                padding: 10px 8px;
                font-weight: 900;
            }
            QTableWidget::item {
                padding: 6px;
                border-bottom: 1px solid #e8e2cd;
            }
            QStatusBar {
                color: #25231e;
                background-color: #ffe477;
                border-top: 2px solid #25231e;
                font-weight: 600;
            }
            """
        )
        self._setup_ui()

    def _setup_ui(self) -> None:
        central_widget = QWidget(self)
        central_widget.setObjectName("mainPanel")
        main_layout = QVBoxLayout(central_widget)
        main_layout.setContentsMargins(22, 18, 22, 18)
        main_layout.setSpacing(13)

        hero_panel = QWidget()
        hero_panel.setObjectName("heroPanel")
        hero_layout = QVBoxLayout(hero_panel)
        hero_layout.setContentsMargins(20, 16, 20, 16)
        hero_layout.setSpacing(5)
        title = QLabel("KOSPI 200 구성종목")
        title.setObjectName("pageTitle")
        description = QLabel(
            "알록달록한 종목 탐험대 · 시가총액 상위 200개를 찾아보세요!"
        )
        description.setObjectName("pageSubtitle")
        hero_layout.addWidget(title)
        hero_layout.addWidget(description)
        main_layout.addWidget(hero_panel)

        control_card = QWidget()
        control_card.setObjectName("controlCard")
        control_layout = QVBoxLayout(control_card)
        control_layout.setContentsMargins(16, 13, 16, 13)
        control_layout.setSpacing(10)
        control_buttons = QHBoxLayout()
        self.crawl_button = QPushButton("크롤링 시작")
        self.crawl_button.setObjectName("crawlButton")
        self.save_button = QPushButton("엑셀로 저장")
        self.save_button.setObjectName("saveButton")
        self.save_button.setEnabled(False)
        control_buttons.addWidget(self.crawl_button)
        control_buttons.addWidget(self.save_button)
        control_buttons.addStretch()
        control_layout.addLayout(control_buttons)

        progress_layout = QHBoxLayout()
        self.progress_label = QLabel("크롤링을 시작하세요.")
        self.progress_label.setObjectName("progressLabel")
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, TARGET_COUNT)
        self.progress_bar.setValue(0)
        progress_layout.addWidget(self.progress_label)
        progress_layout.addWidget(self.progress_bar, 1)
        control_layout.addLayout(progress_layout)
        main_layout.addWidget(control_card)

        search_card = QWidget()
        search_card.setObjectName("searchCard")
        search_layout = QHBoxLayout(search_card)
        search_layout.setContentsMargins(14, 11, 14, 11)
        search_layout.setSpacing(9)
        search_title = QLabel("종목 찾기")
        search_title.setObjectName("sectionTitle")
        self.search_input = QLineEdit()
        self.search_input.setObjectName("searchInput")
        self.search_input.setClearButtonEnabled(True)
        self.search_input.setPlaceholderText("종목명 또는 종목코드 입력")
        self.clear_search_button = QPushButton("검색 지우기")
        self.clear_search_button.setObjectName("clearSearchButton")
        self.result_count_label = QLabel("검색 결과 0 / 0")
        self.result_count_label.setObjectName("resultCount")
        search_layout.addWidget(search_title)
        search_layout.addWidget(self.search_input, 1)
        search_layout.addWidget(self.clear_search_button)
        search_layout.addWidget(self.result_count_label)
        main_layout.addWidget(search_card)

        table_card = QWidget()
        table_card.setObjectName("tableCard")
        table_layout = QVBoxLayout(table_card)
        table_layout.setContentsMargins(12, 12, 12, 12)
        self.table_widget = QTableWidget(0, 5)
        self.table_widget.setObjectName("constituentTable")
        self.table_widget.setHorizontalHeaderLabels(
            ["종목코드", "종목명", "시가총액(원)", "현재가(원)", "등락률"]
        )
        self.table_widget.setAlternatingRowColors(True)
        self.table_widget.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.table_widget.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        header = self.table_widget.horizontalHeader()
        if header is not None:
            header.setSectionResizeMode(
                1, QHeaderView.ResizeMode.Stretch
            )
            for column in (0, 2, 3, 4):
                header.setSectionResizeMode(
                    column, QHeaderView.ResizeMode.ResizeToContents
                )
        table_layout.addWidget(self.table_widget)
        main_layout.addWidget(table_card, 1)
        self.setCentralWidget(central_widget)
        self.status_bar.showMessage("준비")

        self.crawl_button.clicked.connect(self.start_crawling)
        self.save_button.clicked.connect(self.save_excel)
        self.search_input.textChanged.connect(self.filter_constituents)
        self.clear_search_button.clicked.connect(self.search_input.clear)

    def start_crawling(self) -> None:
        if self.worker is not None and self.worker.isRunning():
            return
        self.constituents.clear()
        self.search_input.clear()
        self.table_widget.setRowCount(0)
        self.result_count_label.setText("검색 결과 0 / 0")
        self.save_button.setEnabled(False)
        self.crawl_button.setEnabled(False)
        self.progress_bar.setValue(0)
        self.progress_label.setText("크롤링 중...")
        self.status_bar.showMessage("TradingView에서 종목 데이터를 가져옵니다.")

        self.worker = CrawlerWorker(self)
        self.worker.progress.connect(self.update_progress)
        self.worker.completed.connect(self.crawling_completed)
        self.worker.failed.connect(self.crawling_failed)
        self.worker.finished.connect(self.worker_finished)
        self.worker.start()

    def update_progress(self, page_number: int, found_count: int) -> None:
        self.progress_bar.setValue(min(found_count, TARGET_COUNT))
        self.progress_label.setText(
            f"{page_number}페이지 조회 · {min(found_count, TARGET_COUNT)}"
            f"/{TARGET_COUNT}개 수집"
        )

    def crawling_completed(
        self,
        constituents: list[Constituent],
        match_count: int,
    ) -> None:
        self.constituents = constituents
        self._display_constituents()
        self.save_button.setEnabled(True)
        self.progress_bar.setValue(TARGET_COUNT)
        self.progress_label.setText(f"{len(constituents)}개 종목 수집 완료")
        if match_count > TARGET_COUNT:
            self.status_bar.showMessage(
                f"스캐너 매칭 {match_count}개 중 시가총액 상위 "
                f"{TARGET_COUNT}개를 표시합니다."
            )
        else:
            self.status_bar.showMessage(f"{len(constituents)}개 종목 수집 완료")

    def crawling_failed(self, message: str) -> None:
        self.progress_label.setText("크롤링에 실패했습니다.")
        self.status_bar.showMessage("데이터를 가져오지 못했습니다.")
        QMessageBox.critical(self, "크롤링 오류", message)

    def worker_finished(self) -> None:
        self.crawl_button.setEnabled(True)
        self.worker = None

    def _display_constituents(self) -> None:
        query = self.search_input.text().strip().casefold()
        visible_constituents = [
            constituent
            for constituent in self.constituents
            if not query
            or query in constituent.name.casefold()
            or query in constituent.code.casefold()
        ]
        self.table_widget.setRowCount(len(visible_constituents))
        for row, constituent in enumerate(visible_constituents):
            values = (
                constituent.code,
                constituent.name,
                f"{constituent.market_cap:,.0f}",
                f"{constituent.price:,.0f}",
                f"{constituent.change:+.2f}%",
            )
            for column, value in enumerate(values):
                item = QTableWidgetItem(value)
                if column in (0, 2, 3, 4):
                    item.setTextAlignment(
                        Qt.AlignmentFlag.AlignRight
                        | Qt.AlignmentFlag.AlignVCenter
                    )
                self.table_widget.setItem(row, column, item)
        self.result_count_label.setText(
            f"검색 결과 {len(visible_constituents)} / {len(self.constituents)}"
        )
        if self.constituents:
            self.status_bar.showMessage(
                f"검색 결과 {len(visible_constituents)}개 / "
                f"전체 {len(self.constituents)}개 · 엑셀 저장은 전체 목록"
            )

    def filter_constituents(self) -> None:
        self._display_constituents()

    def save_excel(self) -> None:
        if not self.constituents:
            QMessageBox.information(
                self, "저장할 데이터 없음", "먼저 종목 크롤링을 실행하세요."
            )
            return
        output_path, _ = QFileDialog.getSaveFileName(
            self,
            "엑셀 파일 저장",
            str(OUTPUT_FILE),
            "Excel 파일 (*.xlsx)",
        )
        if not output_path:
            return
        path = Path(output_path)
        if path.suffix.lower() != ".xlsx":
            path = path.with_suffix(".xlsx")
        try:
            save_to_excel(self.constituents, path)
        except (OSError, PermissionError, ValueError) as error:
            QMessageBox.critical(
                self, "엑셀 저장 오류", f"파일 저장에 실패했습니다.\n{error}"
            )
            return
        self.status_bar.showMessage(f"엑셀 저장 완료: {path}", 5000)
        QMessageBox.information(
            self, "저장 완료", f"{len(self.constituents)}개 종목을 저장했습니다.\n{path}"
        )


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = KOSPI200Window()
    window.show()
    sys.exit(app.exec())
